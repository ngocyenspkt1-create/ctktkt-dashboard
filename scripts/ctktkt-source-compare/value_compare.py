"""So sánh GIÁ TRỊ từng ô giữa file Chỉ tiêu gốc và file xuất từ web (cả hai đã được Excel tính lại).
Dùng: python value_compare.py goc.xlsx web.xlsx [ket_qua.json]"""
import openpyxl, math, json, re, sys
from openpyxl.utils import get_column_letter as L

GOC, WEB = sys.argv[1], sys.argv[2]
Gf, Gv = openpyxl.load_workbook(GOC), openpyxl.load_workbook(GOC, data_only=True)
Wf, Wv = openpyxl.load_workbook(WEB), openpyxl.load_workbook(WEB, data_only=True)

# Column blocks and the column holding each block's row labels.
BLOCKS = [(1, 11, [2, 1]), (12, 18, [12]), (21, 29, [22, 21]), (31, 38, [32, 31]), (39, 49, [39])]

def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower() if isinstance(s, str) else None

def labels(ws, cols, maxr):
    out = {}
    for r in range(1, maxr + 1):
        for c in cols:
            v = norm(ws.cell(r, c).value)
            if v and not re.fullmatch(r"[\d.,\-]+", v):
                out[r] = v
                break
    return out

def row_map(gl, wl, maxr):
    # For each goc row, find web row with the same label within +-15 rows (closest), else inherit last offset.
    wl_by = {}
    for r, v in wl.items():
        wl_by.setdefault(v, []).append(r)
    mapping, off = {}, 0
    for r in range(1, maxr + 1):
        if r in gl and gl[r] in wl_by:
            cands = [w for w in wl_by[gl[r]] if abs(w - r) <= 15]
            if cands:
                w = min(cands, key=lambda w: abs(w - (r + off)))
                off = w - r
        mapping[r] = r + off
    return mapping

def num(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None

def same(a, b):
    na, nb = num(a), num(b)
    if na is not None and nb is not None:
        return math.isclose(na, nb, rel_tol=1e-7, abs_tol=1e-7)
    ea, eb = a in (None, ""), b in (None, "")
    if ea or eb:
        return ea and eb or (ea and nb == 0) or (eb and na == 0)
    return norm(a) == norm(b)

result = {}
sheets = ["d-1"] + [s for s in (f"{i:02d}" for i in range(1, 32)) if s in Gf.sheetnames and s in Wf.sheetnames and isinstance(Gf[s]["AB8"].value, (int, float))] + ["Tổng hợp tháng"]
for s in sheets:
    gf, gv, wf, wv = Gf[s], Gv[s], Wf[s], Wv[s]
    maxr = max(gf.max_row, wf.max_row)
    diffs = []
    for c1, c2, lc in BLOCKS:
        mp = row_map(labels(gf, lc, maxr), labels(wf, lc, maxr), maxr)
        for r in range(1, gf.max_row + 1):
            for c in range(c1, min(c2, gf.max_column) + 1):
                a = gv.cell(r, c).value
                fa = gf.cell(r, c).value
                if a in (None, "") and not (isinstance(fa, str) and fa.startswith("=")):
                    continue
                wr = mp[r]
                b = wv.cell(wr, c).value if wr >= 1 else None
                if same(a, b):
                    continue
                fb = wf.cell(wr, c).value if wr >= 1 else None
                diffs.append({
                    "g": f"{L(c)}{r}", "w": f"{L(c)}{wr}", "gv": str(a)[:40], "wv": str(b)[:40],
                    "gf": str(fa)[:80] if isinstance(fa, str) and fa.startswith("=") else "",
                    "wf": str(fb)[:80] if isinstance(fb, str) and str(fb).startswith("=") else "",
                    "lab": (labels(gf, lc, r).get(r) or "")[:45],
                })
    result[s] = diffs
    print(s, len(diffs))
json.dump(result, open(sys.argv[3] if len(sys.argv) > 3 else "fullcmp.json", "w", encoding="utf-8"), ensure_ascii=False)
