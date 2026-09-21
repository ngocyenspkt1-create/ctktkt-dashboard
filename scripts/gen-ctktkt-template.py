from __future__ import annotations

import base64
import io
import json
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.dimensions import DimensionHolder

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT.parent / ".tmp-inspect" / "ctktkt-17-09" / "CHI_TIEU_KTKT_17.09.2026_converted.xlsx"
REFERENCE_SHEET = sys.argv[2] if len(sys.argv) > 2 else "03"
TEMPLATE_TARGET = ROOT / "lib" / "ctktkt-template.generated.ts"
FIELDS_TARGET = ROOT / "lib" / "ctktkt-fields.generated.ts"
CANDIDATES_SOURCE = ROOT / "assets" / "ctktkt-input-candidates.json"

wb = load_workbook(SOURCE, data_only=False, read_only=False)
daily_sheets = [f"{day:02d}" for day in range(1, 32)]
if REFERENCE_SHEET not in wb.sheetnames:
    raise RuntimeError(f"missing reference sheet {REFERENCE_SHEET}")
ws = wb[REFERENCE_SHEET]
candidate_cells = json.loads(CANDIDATES_SOURCE.read_text(encoding="utf-8"))


def section(row: int) -> tuple[str, str]:
    if row <= 20:
        return "power_meters", "Công suất và công tơ chính"
    if row <= 39:
        return "fuel_meters", "Công tơ than, dầu và điện tự dùng"
    if row <= 64:
        return "water_oil", "Nước, dầu HFO và hơi"
    if row <= 79:
        return "environment", "NH3, nước demin và thời gian vận hành"
    if row <= 119:
        return "coal_quality", "Than trộn, độ ẩm và nhiệt trị"
    if row <= 154:
        return "meter_summary", "Bảng tổng hợp công tơ"
    if row <= 176:
        return "pmis_daily", "Đối chiếu PMIS theo ngày"
    return "pmis_02pd", "Báo cáo PMIS 02-PĐ"


def text(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split())


def label_for(cell) -> str:
    col, row = cell.column, cell.row
    anchors = []
    if 13 <= col <= 18:
        anchors.append(ws.cell(row, 12).value)
        anchors.append(ws.cell(2, col).value)
    elif 23 <= col <= 28:
        anchors.append(ws.cell(row, 22).value)
        anchors.append(ws.cell(7 if row <= 39 else 40, col).value)
    elif 33 <= col <= 38:
        anchors.append(ws.cell(row, 32).value)
        anchors.append(ws.cell(7 if row <= 39 else 40, col).value)
    elif row == 181:
        anchors.extend([ws.cell(179, col).value, ws.cell(178, col).value])
    elif row in (157, 158):
        anchors.extend([ws.cell(row, 3).value, ws.cell(156, col).value])
    else:
        for candidate_col in range(col - 1, max(0, col - 12), -1):
            value = ws.cell(row, candidate_col).value
            if isinstance(value, str) and not value.startswith("="):
                anchors.append(value)
                break
        for candidate_row in range(row - 1, max(0, row - 5), -1):
            value = ws.cell(candidate_row, col).value
            if isinstance(value, str) and not value.startswith("="):
                anchors.append(value)
                break
    parts = [text(value) for value in anchors if text(value)]
    return " · ".join(dict.fromkeys(parts)) or f"Ô {cell.coordinate}"


fields = []
reference_input_cells = []
for coordinate in candidate_cells:
    cell = ws[coordinate]
    key, section_label = section(cell.row)
    fields.append({
        "cell": coordinate,
        "section": key,
        "sectionLabel": section_label,
        "label": label_for(cell),
        "row": cell.row,
        "column": cell.column,
    })
    if cell.value is not None and not (isinstance(cell.value, str) and cell.value.startswith("=")):
        reference_input_cells.append(coordinate)

fields.sort(key=lambda item: (item["row"], item["column"]))
reference_input_cells.sort(key=lambda coordinate: (ws[coordinate].row, ws[coordinate].column))


def quoted_sheet_name(name: str) -> str:
    return "'" + name.replace("'", "''") + "'!"


def rebased_value(value: object, target_name: str, previous_name: str) -> object:
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    previous_marker = "__CTKTKT_PREVIOUS_SHEET__!"
    current_marker = "__CTKTKT_CURRENT_SHEET__!"
    formula = value.replace("'02'!", previous_marker).replace("02!", previous_marker)
    formula = formula.replace("'03'!", current_marker).replace("03!", current_marker)
    return formula.replace(previous_marker, quoted_sheet_name(previous_name)).replace(current_marker, quoted_sheet_name(target_name))


def clone_reference_layout(reference, target, target_name: str, previous_name: str) -> None:
    for merged_range in list(target.merged_cells.ranges):
        target.unmerge_cells(str(merged_range))
    max_row = max(reference.max_row, target.max_row)
    max_column = max(reference.max_column, target.max_column)
    for row in range(1, max_row + 1):
        for column in range(1, max_column + 1):
            source_cell = reference.cell(row, column)
            target_cell = target.cell(row, column)
            target_cell.value = rebased_value(source_cell.value, target_name, previous_name)
            target_cell._style = source_cell._style
            target_cell.hyperlink = copy(source_cell.hyperlink)
            target_cell.comment = None
    for merged_range in reference.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    target.row_dimensions = DimensionHolder(worksheet=target)
    for key, dimension in reference.row_dimensions.items():
        cloned = copy(dimension)
        cloned.worksheet = target
        target.row_dimensions[key] = cloned
    target.column_dimensions = DimensionHolder(worksheet=target)
    for key, dimension in reference.column_dimensions.items():
        cloned = copy(dimension)
        cloned.worksheet = target
        target.column_dimensions[key] = cloned
    target.sheet_format = copy(reference.sheet_format)
    target.sheet_properties = copy(reference.sheet_properties)
    target.page_margins = copy(reference.page_margins)
    target.page_setup = copy(reference.page_setup)
    target.print_options = copy(reference.print_options)
    target.views = copy(reference.views)
    target.freeze_panes = reference.freeze_panes
    target.auto_filter = copy(reference.auto_filter)
    target.print_area = reference.print_area
    target.print_title_cols = reference.print_title_cols
    target.print_title_rows = reference.print_title_rows
    target.sheet_state = reference.sheet_state


for day, name in enumerate(daily_sheets, start=1):
    if name == REFERENCE_SHEET:
        continue
    previous_name = "d-1" if day == 1 else f"{day - 1:02d}"
    clone_reference_layout(ws, wb[name], name, previous_name)

# Never commit the user's real operating figures. Build a clean template that
# retains formulas/styles/layout but removes every identified manual field and
# other literal values from operational input blocks. Saving with openpyxl also
# removes cached formula results while keeping the formulas themselves.
clear_ranges = [
    "M3:R20", "W8:AB28", "AG8:AL28", "M41:R46", "M49:R64",
    "W54:AB54", "AG54:AL54", "W68:Z74", "N69:Q75", "M81:O82",
    "W86:W89", "AJ87:AO92", "J157:L158", "C181:T181",
]
for name in ["d-1", *daily_sheets]:
    sheet = wb[name]
    for field in fields:
        target_cell = sheet[field["cell"]]
        if not isinstance(target_cell, MergedCell):
            target_cell.value = None
    for range_ref in clear_ranges:
        for row in sheet[range_ref]:
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    cell.value = None
    for target in ["Z57", "AJ57", "AE83"]:
        sheet[target].value = "Ngày vận hành"
    for target in ["N68", "O68"]:
        sheet[target].value = "Mức bồn theo ngày vận hành"

# Comments in the source are tied to legacy VML drawing parts. They may contain
# real operational notes and openpyxl cannot safely rewrite those VML links for
# ExcelJS, so the distributable template intentionally omits comments.
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell) and cell.comment is not None:
                cell.comment = None

month_sheet = wb["Tổng hợp tháng"]
for row in month_sheet.iter_rows(min_row=3, max_row=26, min_col=2, max_col=40):
    for cell in row:
        if isinstance(cell.value, (int, float)):
            cell.value = None

buffer = io.BytesIO()
wb.save(buffer)
sanitized_bytes = buffer.getvalue()
cached_book = load_workbook(io.BytesIO(sanitized_bytes), data_only=True, read_only=False)
formula_cache_count = 0
for name in ["d-1", *daily_sheets, "Tổng hợp tháng"]:
    formula_sheet = wb[name]
    cached_sheet = cached_book[name]
    for row in formula_sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=") and cached_sheet[cell.coordinate].value is not None:
                formula_cache_count += 1
if formula_cache_count:
    raise RuntimeError(f"sanitized template still contains {formula_cache_count} cached formula results")
encoded = base64.b64encode(sanitized_bytes).decode("ascii")
template_target = (
    "// Generated from the user-supplied workbook. Do not edit by hand.\n"
    f"export const CTKTKT_TEMPLATE_BASE64 = {json.dumps(encoded)};\n"
)
fields_target = (
    "// Generated from the user-supplied workbook. Do not edit by hand.\n"
    f"export const CTKTKT_INPUT_FIELDS = {json.dumps(fields, ensure_ascii=False, separators=(',', ':'))} as const;\n"
    f"export const CTKTKT_DAY03_INPUT_CELLS = {json.dumps(reference_input_cells, ensure_ascii=False, separators=(',', ':'))} as const;\n"
)
TEMPLATE_TARGET.write_text(template_target, encoding="utf-8")
FIELDS_TARGET.write_text(fields_target, encoding="utf-8")
print(
    f"generated {TEMPLATE_TARGET} and {FIELDS_TARGET} from sheet {REFERENCE_SHEET} "
    f"with {len(fields)} fields, {len(reference_input_cells)} reference inputs and {len(encoded)} base64 characters"
)
