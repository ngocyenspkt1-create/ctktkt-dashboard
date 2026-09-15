"""Import manual-entry cells from monthly T01..T12 sheets into the local app."""

from __future__ import annotations

import argparse
import calendar
import json
import re
import urllib.request
import urllib.error
from pathlib import Path

import openpyxl


FIELD_CODES = (
    "B", "C", "F", "H", "I", "L", "X", "AE", "AF", "AJ", "AR", "AT",
    "CJ", "CX", "BN", "BQ", "BR", "CN", "BZ", "CA", "CC", "CD", "CE",
    "CF", "CM", "CQ", "CR", "CS", "CT", "CU", "CV", "CW",
)


def text_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".12g")
    return str(value).strip()


def request_json(url: str, payload: dict | None = None) -> dict:
    data = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"} if data else {}, method="POST" if data else "GET")
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {error.code}: {detail}") from error


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--base-url", default="http://localhost:5173")
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    formulas = openpyxl.load_workbook(args.workbook, read_only=False, data_only=False, keep_links=False)
    values = openpyxl.load_workbook(args.workbook, read_only=False, data_only=True, keep_links=False)
    sheets = sorted((name for name in formulas.sheetnames if re.fullmatch(r"T\d{2}", name)), key=lambda name: int(name[1:]))
    report: list[dict] = []

    for sheet_name in sheets:
        month = int(sheet_name[1:])
        if not 1 <= month <= 12:
            continue
        ws_formula, ws_value = formulas[sheet_name], values[sheet_name]
        days = calendar.monthrange(args.year, month)[1]
        period = f"{args.year}-{month:02d}"
        entries: list[dict] = []
        uncached: list[str] = []
        invalid_numeric: list[dict] = []

        for day in range(1, days + 1):
            row = day + 5
            for code in FIELD_CODES:
                raw = ws_formula[f"{code}{row}"].value
                value = ws_value[f"{code}{row}"].value if isinstance(raw, str) and raw.startswith("=") else raw
                if isinstance(raw, str) and raw.startswith("=") and value is None:
                    uncached.append(f"{code}{row}")
                    continue
                serialized = text_value(value)
                if serialized == "":
                    continue
                if code != "CW":
                    normalized = serialized.replace(" ", "")
                    if re.fullmatch(r"-?\d+(?:[.,]\d+)?", normalized):
                        serialized = normalized
                    elif not (code == "CX" and serialized.casefold() == "không chỉnh"):
                        invalid_numeric.append({"cell": f"{code}{row}", "value": serialized})
                        continue
                entries.append({
                    "operatingDate": f"{period}-{day:02d}",
                    "fieldCode": code,
                    "value": serialized,
                    "note": "",
                })

        item = {"sheet": sheet_name, "period": period, "sourceEntries": len(entries), "uncachedFormulas": uncached, "invalidNumeric": invalid_numeric}
        if args.apply:
            response = request_json(f"{args.base_url}/api/daily-inputs", {"period": period, "entries": entries})
            loaded = request_json(f"{args.base_url}/api/daily-inputs?period={period}").get("entries", [])
            item.update({"saved": response.get("saved"), "loadedEntries": len(loaded), "verified": len(loaded) == len(entries)})
        report.append(item)

    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
